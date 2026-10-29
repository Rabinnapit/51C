"""Package the offline file as a Word document for email (Army email blocks .html attachments).

  python tools/make_docx.py        (run tools/build.py first; needs python-docx)

Writes to dist/docx/:
  1 - READ ME FIRST - How to open B2B Exam Prep.docx     step-by-step instructions
  2 - B2B-Exam-Prep.html.docx                            the app's code, one line per paragraph
The code document is read back after saving and compared with the HTML character by character; the script stops
if anything differs, so a damaged package never goes out. (Adapted from Board Ready's make_docx.py.)
"""
import hashlib
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
DIST = ROOT / "dist"
OUT = DIST / "docx"
OUT.mkdir(parents=True, exist_ok=True)
HTML = "B2B-Exam-Prep.html"
CODE_DOC = "2 - B2B-Exam-Prep.html.docx"
README_DOC = "1 - READ ME FIRST - How to open B2B Exam Prep.docx"


def no_proof(run):
    run._r.get_or_add_rPr().append(OxmlElement("w:noProof"))


def quiet(doc):
    for tag in ("w:hideSpellingErrors", "w:hideGrammaticalErrors"):
        doc.settings.element.append(OxmlElement(tag))


def code_doc(src: Path, out: Path):
    text = src.read_text(encoding="utf-8")
    doc = Document(); quiet(doc)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(0.5)
    st = doc.styles["Normal"]
    st.font.name = "Consolas"; st.element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas"); st.font.size = Pt(7)
    pf = st.paragraph_format; pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)
    lines = text.split("\n")
    for line in lines:
        p = doc.add_paragraph()
        if line:
            no_proof(p.add_run(line))
    doc.save(out)
    back = "\n".join(p.text for p in Document(out).paragraphs)
    if back != text:
        n = next((i for i in range(min(len(back), len(text))) if back[i] != text[i]), min(len(back), len(text)))
        raise SystemExit(f"ROUND-TRIP MISMATCH at character {n}: {back[n:n+40]!r} vs {text[n:n+40]!r}")
    return len(lines), len(text.encode("utf-8")), hashlib.sha256(text.encode("utf-8")).hexdigest()[:12]


def readme(out: Path, lines: int, utf8: int, site_url: str):
    doc = Document(); quiet(doc)
    sec = doc.sections[0]; sec.left_margin = sec.right_margin = Inches(1)
    doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("B2B Exam Prep: how to open it from these Word documents", level=1)
    doc.add_paragraph("Army email blocks .html attachments, so B2B Exam Prep travels inside a Word document as plain text. Word turns it back "
                      "into a working file in about a minute. You need Microsoft Word and the Edge or Chrome browser. No administrator rights "
                      "and no internet connection are needed.")
    if site_url:
        doc.add_paragraph(f"Easier option: if your network allows it, just open the website instead: {site_url}  The website always has the "
                          "newest version and can be installed as an app.")
    doc.add_heading("Turn the code document back into a file", level=2)
    for s in [
        "Open “2 - B2B-Exam-Prep.html.docx”. It looks like hundreds of pages of code; that is expected. If a yellow Protected View bar appears, click Enable Editing. Do not change any text.",
        "Choose File, then Save As, then Browse. If you only see Save a Copy, use that.",
        "Pick a folder you can find again, for example a new folder named B2B Exam Prep in Documents or OneDrive.",
        "Set “Save as type” to “Plain Text (*.txt)”.",
        "In “File name”, type the name with the quotation marks: \"B2B-Exam-Prep.html\". Then click Save.",
        "A File Conversion window opens. Choose “Other encoding”, then “Unicode (UTF-8)”. Leave “Insert line breaks” and “Allow character substitution” unchecked. Click OK, then close Word.",
        "Open the folder and double-click B2B-Exam-Prep.html. It opens in your browser and works with no network.",
    ]:
        doc.add_paragraph(s, style="List Number")
    word_bytes = utf8 + (lines - 1) + 3 + 2
    doc.add_heading("Check the result", level=2)
    doc.add_paragraph(f"Right-click the saved file and choose Properties. It should be about {word_bytes:,} bytes "
                      f"({word_bytes / 1024 / 1024:.1f} MB). A difference of a few bytes is fine; a much smaller file means the save was cut short.")
    doc.add_heading("If something goes wrong", level=2)
    for s in [
        "The file name ends in .html.txt: in File Explorer choose View, Show, File name extensions, then rename it so it ends in .html.",
        "The browser shows a page of code instead of the guide: the file still ends in .txt, or the encoding was not UTF-8. Repeat steps 4 to 6.",
    ]:
        doc.add_paragraph(s, style="List Bullet")
    doc.add_heading("Keep your progress", level=2)
    doc.add_paragraph("Your ticks, scores and flashcards are saved in the browser, and some government computers clear that at restart. "
                      "Click Save progress (top of the page) and choose this same B2B-Exam-Prep.html file; replace it when asked. Your progress is "
                      "then stored inside the file. If the browser downloads a copy instead, move that copy into your folder and replace the old file.")
    doc.add_heading("Sharing", level=2)
    doc.add_paragraph("Once saved, the file can go on Teams or SharePoint or be emailed inside the Army network, and people just open it. "
                      "A file with progress saved inside carries that progress, so share a fresh copy, not your own.")
    doc.add_paragraph("Unofficial training aid. Not endorsed by the U.S. Army, the Department of War or WarU. Verify rules and numbers with your "
                      "Contracting Master Gunner (KMG) and contracting policy office before relying on them.")
    doc.save(out)


if __name__ == "__main__":
    import json, re
    src = DIST / HTML
    if not src.exists():
        raise SystemExit("dist/B2B-Exam-Prep.html is missing; run tools/build.py first")
    lines, utf8, digest = code_doc(src, OUT / CODE_DOC)
    m = re.search(r'"site_url":"([^"]*)"', src.read_text(encoding="utf-8"))
    readme(OUT / README_DOC, lines, utf8, m.group(1) if m else "")
    print(f"{CODE_DOC}: {lines:,} lines, {utf8:,} bytes, sha {digest}, round-trip OK")
    print(f"{README_DOC}: written")
    # the website offers the email package for download too (docs/ is what GitHub Pages publishes)
    docs = ROOT / "docs"
    if docs.exists():
        import shutil
        dl = docs / "downloads"
        dl.mkdir(exist_ok=True)
        for name, target in ((README_DOC, "B2B-Exam-Prep-READ-ME-FIRST.docx"), (CODE_DOC, "B2B-Exam-Prep.html.docx")):
            shutil.copyfile(OUT / name, dl / target)
        print(f"copied both documents to {dl}")
