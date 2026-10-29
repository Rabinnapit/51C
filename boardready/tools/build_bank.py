"""Build the normalized CORB question bank from the PAM xlsx, the older CSV, and the study decks.
Outputs:
  app/src/data/questions.json      -- the bank (list of question objects)
  app/content/bank_changelog.md    -- every merge / re-topic / cleanup decision
  app/content/citations.json       -- unique citations found in answers, with counts and question ids
  app/content/open_questions.json  -- questions with no formal answer (for drafting agents)
"""
import os, re, csv, json, html, difflib
from pathlib import Path
from collections import Counter, defaultdict
import openpyxl, pptx

APP = Path(__file__).resolve().parents[1]
def _sources_root():
    """The folder that holds corb/WarrantQA.csv and CORB2/: CORB_SOURCES, or the nearest parent of the app that has them."""
    env = os.environ.get("CORB_SOURCES")
    if env:
        return Path(env)
    for parent in APP.parents:
        if (parent / "corb" / "WarrantQA.csv").exists() or (parent / "CORB2").exists():
            return parent
    return APP.parents[1]


ROOT = _sources_root()          # .../CORB (holds corb/ and CORB2/)
XLSX = ROOT / "CORB2" / "Warrant Board Study Group - Question and Answer List.xlsx"
CSV = ROOT / "corb" / "WarrantQA.csv"
DECKS = sorted((ROOT / "CORB2").glob("Warrant Board Study Group*.pptx"))
OUT_DATA = ROOT / "app" / "src" / "data"
OUT_CONTENT = ROOT / "app" / "content"
OUT_DATA.mkdir(parents=True, exist_ok=True)
OUT_CONTENT.mkdir(parents=True, exist_ok=True)

log = []


def L(msg):
    log.append(msg)


# ---------- text cleaning ----------
def clean(t):
    if t is None:
        return ""
    t = str(t)
    t = t.replace("​", "").replace("﻿", "")
    t = re.sub(r"(?i)<br\s*/?>", "\n", t)
    t = re.sub(r"(?i)</p>|</li>|</div>|</h\d>", "\n", t)
    t = re.sub(r"(?i)<li[^>]*>", "- ", t)
    t = re.sub(r"<[^>]+>", "", t)
    t = html.unescape(t).replace("\xa0", " ")
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n\s*\n+", "\n\n", t)
    return t.strip()


def norm_key(q):
    q = clean(q).lower()
    q = re.sub(r"[^a-z0-9 ]+", " ", q)
    return re.sub(r"\s+", " ", q).strip()


OPEN_MARK = "intentionally not populated"


def is_open(ans):
    return (not ans) or (OPEN_MARK in ans.lower()) or len(ans.strip()) < 8


def norm_classes(v):
    v = clean(v)
    parts = None
    if v.startswith("["):
        try:
            parts = [str(x) for x in json.loads(v)]
        except Exception:
            parts = None
    if parts is None:
        parts = re.split(r";#|,", v)
    out = []
    for part in parts:
        p = part.strip().replace("Class ", "").replace("Class", "").strip()
        if not p:
            continue
        if p in ("I", "II", "III", "IV"):
            out.append(p)
        elif p.lower().startswith("agree"):
            out.append("Agreements")
        elif p.lower().startswith("grant"):
            out.append("Grants")
        else:
            out.append(p)
    order = {"I": 0, "II": 1, "III": 2, "IV": 3, "Agreements": 4, "Grants": 5}
    return sorted(dict.fromkeys(out), key=lambda x: order.get(x, 9))


def split_tags(v, sep=";#"):
    v = clean(v)
    if v.startswith("["):
        try:
            return [t.strip() for t in json.loads(v) if t.strip()]
        except Exception:
            pass
    return [t.strip() for t in v.split(sep) if t.strip()]


# ---------- domain mapping (topic -> navigation domain) ----------
DOMAIN_RULES = [
    (r"korea|\blcs\b|sofa|invited contractor|usfk", "Korea, LCS & SOFA"),
    (r"abg|bauamt|indirect", "Germany ABG-75"),
    (r"fiscal|funding|bona fide|anti-?deficiency|appropriat|obligat", "Fiscal Law & Funding"),
    (r"grant", "Grants & Agreements"),
    (r"agreement|ota|other transaction|prototype", "Grants & Agreements"),
    (r"protest|cica", "Protests, Claims & Disputes"),
    (r"claim|dispute|recourse", "Protests, Claims & Disputes"),
    (r"terminat|cure|default|stop work|excusable|liquidated|latent", "Terminations & Remedies"),
    (r"source selection|evaluation|trade-?off|lpta|debrief|competitive range|best value", "Source Selection & Evaluation"),
    (r"pricing|price|cost|cas\b|unbalanced|ige|independent government|progress payment|performance based payment|payment", "Pricing, Cost & Payment"),
    (r"competition|sole source|market research|synopsis|publiciz|brand name|consolidat|strategic sourcing|mandatory source|unsolicited", "Competition & Market Research"),
    (r"acquisition planning|requirements|contract type|options?\b|ordering|idiq|delivery order|task order|blanket|bpa|simplified|commercial|quotes|services|personal|performance based|inherently|d&f|determination", "Acquisition Planning & Contract Types"),
    (r"administration|modification|cor\b|cpars|closeout|contract files|fpds|inspection|acceptance|order of precedence|subcontract|post award|qasp|change", "Contract Administration"),
    (r"responsib|sam\b|epls|exclu|debar|suspen", "Responsibility & Exclusions"),
    (r"property|gfp|contractor acquired|berry|security|at/opsec|defense base act|oci|conflict", "Property, Security & Special Provisions"),
    (r"ethic|authority|contracting officer|binding|leadership|mentor|training|management|guidance|general|unauthorized|ratif", "Authority, Ethics & Leadership"),
    (r"small business|socio", "Small Business & Socioeconomic"),
]


def domain_for(topic, tags, question):
    blob = (topic + " " + " ".join(tags)).lower()
    for pat, dom in DOMAIN_RULES:
        if re.search(pat, topic.lower()):
            return dom
    for pat, dom in DOMAIN_RULES:
        if re.search(pat, blob):
            return dom
    ql = question.lower()
    for pat, dom in DOMAIN_RULES:
        if re.search(pat, ql):
            return dom
    return "General Contracting"


def kind_for(topic, tags, question):
    t = (topic + " " + " ".join(tags)).lower()
    q = question.lower()
    if "situational" in q or re.search(r"\byou (are|have|receive|received|notice|awarded)\b|your (customer|contractor|cor|specialist|contract)\b", q):
        return "scenario"
    if re.search(r"mentor|leadership|team|workload|management", t):
        return "leadership"
    return "knowledge"


# ---------- load xlsx ----------
wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
xl = []
for i, r in enumerate(rows[1:], start=2):
    if not any(r):
        continue
    topic, q, a, tags, warr, sco = (clean(r[0]), clean(r[1]), clean(r[2]), r[3], r[4], clean(r[5]))
    if not q:
        continue
    xl.append(dict(src="pam-xlsx-2026-01", row=i, topic=topic, question=q, answer=a, tags=split_tags(tags), classes=norm_classes(warr), unit=sco or "ACC-All"))

# ---------- load csv ----------
cs = []
with open(CSV, encoding="utf-8-sig") as f:
    for i, r in enumerate(csv.DictReader(f), start=2):
        q = clean(r["Question"])
        a = clean(r["Answer"])
        if not q:
            continue
        cs.append(dict(src="pam-csv-legacy", row=i, topic=clean(r["Topic"]), question=q, answer=a, tags=split_tags(r["Topics Applicable to –"]), classes=norm_classes(r["Warrant"]), unit=clean(r["SCO Location"]) or "ACC-All"))

# ---------- load decks ----------
deck_answers = {}     # xlsx row -> list of (deckdate, answer text)
deck_scenarios = []
MONTHS = {"JAN": 1, "FEB": 2, "MAR": 3, "APR": 4, "MAY": 5, "JUN": 6, "JUL": 7, "AUG": 8, "SEP": 9, "OCT": 10, "NOV": 11, "DEC": 12}
for p in DECKS:
    pr = pptx.Presentation(p)
    m = re.search(r"(\d{1,2})([A-Z]{3})(\d{4})", p.name)
    date = f"{m.group(3)}-{MONTHS[m.group(2)]:02d}-{int(m.group(1)):02d}"
    slides = []
    for s in pr.slides:
        texts = []
        for sh in s.shapes:
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs)
                    if t.strip():
                        texts.append(t.strip())
        slides.append(texts)
    for idx, texts in enumerate(slides):
        if not texts:
            continue
        head = texts[0]
        nxt = slides[idx + 1] if idx + 1 < len(slides) else []
        nxt_is_answer = bool(nxt) and nxt[0].lower().startswith("answer")
        if head.lower().startswith("scenario-based question"):
            qtext = "\n".join(texts[1:])
            atext = "\n".join(nxt[1:]) if nxt_is_answer else ""
            deck_scenarios.append(dict(src=f"deck-{date}", slide=idx + 1, question=qtext, answer=atext))
        mm = re.match(r"Q&A List Question\s*[–-]\s*Line\s*([\d/]+)", head)
        if mm and nxt_is_answer:
            for ln in mm.group(1).split("/"):
                deck_answers.setdefault(int(ln), []).append((date, "\n".join(nxt[1:])))

# ---------- merge ----------
bank = []
index = {}


def add_source(rec, s):
    rec["sources"].append({"src": s["src"], "row": s["row"]})
    if s["unit"] not in rec["units"]:
        rec["units"].append(s["unit"])
    for t in s["tags"]:
        if t not in rec["tags"]:
            rec["tags"].append(t)
    for c in s["classes"]:
        if c not in rec["classes"]:
            rec["classes"].append(c)
    if s["topic"] and s["topic"] not in rec["alt_topics"] and s["topic"] != rec["topic"]:
        rec["alt_topics"].append(s["topic"])
    if is_open(rec["answer"]) and not is_open(s["answer"]):
        rec["answer"] = s["answer"]
        rec["answer_source"] = s["src"]
    elif not is_open(s["answer"]) and s["answer"] != rec["answer"]:
        rec["alt_answers"].append({"src": s["src"], "text": s["answer"]})


def new_record(s):
    return dict(question=s["question"], answer=s["answer"], answer_source=s["src"], topic=s["topic"] or "General",
                alt_topics=[], tags=list(s["tags"]), classes=list(s["classes"]), units=[s["unit"]],
                sources=[{"src": s["src"], "row": s["row"]}], alt_answers=[], deck_notes=[])


def find_near(key):
    if key in index:
        return index[key], 1.0
    best, score = None, 0
    for k, rec in index.items():
        r = difflib.SequenceMatcher(None, key, k).ratio()
        if r > score:
            best, score = rec, r
    return (best, score) if score >= 0.90 else (None, score)


for s in xl + cs:
    key = norm_key(s["question"])
    rec, score = find_near(key)
    if rec:
        if score < 1.0:
            L(f"MERGE near-duplicate ({score:.2f}): {s['src']} row {s['row']} -> existing '{rec['question'][:70]}...'")
        else:
            L(f"MERGE exact duplicate: {s['src']} row {s['row']} -> existing '{rec['question'][:70]}...'")
        add_source(rec, s)
    else:
        rec = new_record(s)
        index[key] = rec
        bank.append(rec)

# deck enrichment by xlsx row
row_to_rec = {}
for rec in bank:
    for src in rec["sources"]:
        if src["src"] == "pam-xlsx-2026-01":
            row_to_rec[src["row"]] = rec
for ln, items in deck_answers.items():
    rec = row_to_rec.get(ln)
    if not rec:
        L(f"DECK: line {ln} not found in xlsx")
        continue
    for date, text in items:
        rec["deck_notes"].append({"src": f"deck-{date}", "text": text})
        if is_open(rec["answer"]):
            rec["answer"] = text
            rec["answer_source"] = f"deck-{date}"
            L(f"DECK: open question row {ln} answered from study deck {date}")
        else:
            L(f"DECK: row {ln} enriched with study-deck answer {date}")

# deck scenarios as new questions
SCEN_META = {
    "power conditioners": ("Competition & Market Research", "Sole Source / Brand Name", ["I", "II"], []),
    "LCS purchase request": ("Korea, LCS & SOFA", "LCS Procurement", ["II", "III"], ["Korea", "LCS Funding"]),
    "water line bursts": ("Authority, Ethics & Leadership", "Unauthorized Commitment", ["I", "II"], ["Korea"]),
    "janitorial services contract": ("Fiscal Law & Funding", "Bona Fide Need", ["I", "II"], ["Fiscal Law / Funding"]),
    "Government Accountability Office": ("Protests, Claims & Disputes", "GAO Protest", ["II"], ["Protests"]),
}
for sc in deck_scenarios:
    dom, topic, classes, tags = ("General Contracting", "Scenario", ["I", "II"], [])
    for k, v in SCEN_META.items():
        if k.lower() in sc["question"].lower():
            dom, topic, classes, tags = v
            break
    rec = dict(question=sc["question"], answer=sc["answer"], answer_source=sc["src"] if sc["answer"] else "", topic=topic, alt_topics=[],
               tags=tags + ["Problem solving situations"], classes=classes, units=["411th CSB"],
               sources=[{"src": sc["src"], "slide": sc["slide"]}], alt_answers=[], deck_notes=[], _domain=dom, _kind="scenario")
    bank.append(rec)
    L(f"DECK: added 411th scenario from {sc['src']} slide {sc['slide']}: '{sc['question'][:70]}...' answer={'yes' if sc['answer'] else 'OPEN'}")

# ---------- cleanup: re-topic + junk tags ----------
RETOPIC = {
    r"limitations of using lcs": "LCS Funding",
    r"both o&m and lcs": "LCS Funding",
    r"logistics cost sharing": "LCS Program",
    r"unique aspects of lcs": "LCS Funding",
    r"invited contractor": "SOFA Invited Contractor Program",
    r"defense base act": "Defense Base Act",
    r"apparent, implied and express": "Contracting Officer Authority",
    r"abg\s*3|abg3|abg7|abg8|abg9|abg-?75|bauamt|kvm|vke|indirect (contracting )?procedure|project specific additional|not paying fees": "Germany ABG-75",
}
for rec in bank:
    if rec["topic"] in ("Acceptance", "", "General"):
        for pat, t in RETOPIC.items():
            if re.search(pat, rec["question"].lower()):
                L(f"RETOPIC: '{rec['question'][:60]}...' '{rec['topic']}' -> '{t}'")
                rec["alt_topics"].append(rec["topic"])
                rec["topic"] = t
                break
    if rec["topic"] == "Acceptance" and not re.search(r"accept", rec["question"].lower()):
        rec["alt_topics"].append("Acceptance")
        rec["topic"] = "General Contracting"
        L(f"RETOPIC: '{rec['question'][:60]}...' 'Acceptance' (SharePoint default) -> 'General Contracting'")
    if rec["topic"] in ("PCO CORB Question (USACE)", "PCO CORB QUESTION"):
        rec["alt_topics"].append(rec["topic"])
        rec["topic"] = "USACE PCO Board Question"
    junk = [t for t in rec["tags"] if re.fullmatch(r"\d+", t)]
    for t in junk:
        rec["tags"].remove(t)
        L(f"TAG: removed junk tag '{t}' from '{rec['question'][:60]}...'")

# ---------- citations ----------
CIT_RE = re.compile(r"\b(DFARS PGI|AFARS PGI|PGI|FAR|DFARS|AFARS)\s*(?:Part|Subpart)?\s*(\d{1,2}(?:\.\d{3,4})?(?:-\d{1,3})?(?:\([a-zA-Z0-9]{1,3}\))*)", re.I)
USC_RE = re.compile(r"\b(\d{1,2})\s*U\.?S\.?C\.?\s*§?\s*(\d{3,5}[a-z]?)", re.I)
CFR_RE = re.compile(r"\b(\d{1,2})\s*C\.?F\.?R\.?\s*(?:§\s*)?(\d{1,4}(?:\.\d+)?)", re.I)
citations = defaultdict(lambda: {"count": 0, "qids": []})

# ---------- finalize ----------


def sort_key(r):
    s0 = r["sources"][0]
    if s0["src"].startswith("pam-xlsx"):
        return s0.get("row", 9999)
    if s0["src"].startswith("pam-csv"):
        return 5000 + s0.get("row", 0)
    return 9000 + s0.get("slide", 0)


bank.sort(key=sort_key)
final = []
for n, rec in enumerate(bank, start=1):
    qid = f"Q{n:04d}"
    dom = rec.pop("_domain", None) or domain_for(rec["topic"], rec["tags"], rec["question"])
    kind = rec.pop("_kind", None) or kind_for(rec["topic"], rec["tags"], rec["question"])
    open_ = is_open(rec["answer"])
    cits = []
    blob = rec["answer"] + " " + " ".join(d["text"] for d in rec["deck_notes"]) + " " + " ".join(a["text"] for a in rec["alt_answers"])
    for m in CIT_RE.finditer(blob):
        c = f"{m.group(1).upper()} {m.group(2)}".replace("DFARS PGI", "PGI")
        cits.append(c)
    for m in USC_RE.finditer(blob):
        cits.append(f"{m.group(1)} U.S.C. {m.group(2)}")
    for m in CFR_RE.finditer(blob):
        cits.append(f"{m.group(1)} CFR {m.group(2)}")
    cits = list(dict.fromkeys(cits))
    for c in cits:
        citations[c]["count"] += 1
        citations[c]["qids"].append(qid)
    units = rec["units"]
    primary_unit = "411th CSB" if "411th CSB" in units else ("ACC-All" if "ACC-All" in units else units[0])
    if rec["topic"] in ("General Contracting", "General", ""):
        rec["topic"] = dom
    final.append({
        "id": qid, "question": rec["question"],
        "answer": "" if open_ else rec["answer"],
        "answer_status": "open" if open_ else "official",
        "answer_source": "" if open_ else rec["answer_source"],
        "deck_notes": rec["deck_notes"], "alt_answers": rec["alt_answers"],
        "domain": dom, "topic": rec["topic"], "alt_topics": rec["alt_topics"], "tags": rec["tags"],
        "classes": rec["classes"] or ["I", "II", "III"], "unit": primary_unit, "units": units,
        "kind": kind, "sources": rec["sources"], "citations": cits,
        "rfo": {"status": "pending", "crosswalk": [], "notes": ""},
        "flags": ["verify-locally"] if primary_unit in ("411th CSB", "409th CSB") else [],
    })

(OUT_DATA / "questions.json").write_text(json.dumps(final, ensure_ascii=False, indent=1), encoding="utf-8")
(OUT_CONTENT / "open_questions.json").write_text(json.dumps([q for q in final if q["answer_status"] == "open"], ensure_ascii=False, indent=1), encoding="utf-8")
cit_sorted = dict(sorted(citations.items(), key=lambda kv: -kv[1]["count"]))
(OUT_CONTENT / "citations.json").write_text(json.dumps(cit_sorted, ensure_ascii=False, indent=1), encoding="utf-8")

stats = Counter(q["answer_status"] for q in final)
dom_c = Counter(q["domain"] for q in final)
unit_c = Counter(q["unit"] for q in final)
kind_c = Counter(q["kind"] for q in final)
md = ["# Question bank build log", "",
      f"Inputs: xlsx rows={len(xl)}, csv rows={len(cs)}, deck scenario questions={len(deck_scenarios)}, deck answers attached={sum(len(v) for v in deck_answers.values())}",
      f"Output: {len(final)} unique questions. answer_status: {dict(stats)}", "",
      "## By domain", *[f"- {k}: {v}" for k, v in dom_c.most_common()], "",
      "## By unit", *[f"- {k}: {v}" for k, v in unit_c.most_common()], "",
      "## By kind", *[f"- {k}: {v}" for k, v in kind_c.most_common()], "",
      f"## Unique citations found: {len(cit_sorted)}", "",
      "## Decisions", *[f"- {m}" for m in log]]
(OUT_CONTENT / "bank_changelog.md").write_text("\n".join(md), encoding="utf-8")
print("\n".join(md[: md.index("## Decisions")]))
print("decisions logged:", len(log))
print("top citations:", list(cit_sorted.items())[:30])
