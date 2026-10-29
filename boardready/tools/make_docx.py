"""Package the built HTML files as Word documents for email transport (Army mail strips .html attachments).

  python tools/make_docx.py

Produces in dist/docx/:
  1 - READ ME FIRST - How to install Board Ready.docx   (instructions)
  2 - BoardReady.html.docx                              (full edition: candidates, board members, facilitators)
  3 - BoardReady-Candidate.html.docx                    (candidate edition: study only)
  4 - BoardReady-Candidate-Offline.html.docx            (candidate edition without the AI connector)
Each code document holds the HTML source, one line per paragraph.
Every code document is round-tripped: the text is read back out of the .docx and compared to the HTML
character by character. The script fails loudly if anything differs.
"""
import sys, hashlib
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

APP = Path(__file__).resolve().parents[1]
DIST = APP / "dist"
OUT = DIST / "docx"
OUT.mkdir(exist_ok=True)

# (built file, Word document, who it is for); the Arcade is inside every file
FILES = [
    ("BoardReady.html", "2 - BoardReady.html.docx", "Full version: candidates, board members and facilitators"),
    ("BoardReady-Candidate.html", "3 - BoardReady-Candidate.html.docx", "Candidate version: study only"),
    ("BoardReady-Candidate-Offline.html", "4 - BoardReady-Candidate-Offline.html.docx", "Candidate version without the AI connector"),
]


def no_proof(run):
    rpr = run._r.get_or_add_rPr()
    el = OxmlElement("w:noProof")
    rpr.append(el)


def quiet_settings(doc):
    st = doc.settings.element
    for tag in ("w:hideSpellingErrors", "w:hideGrammaticalErrors"):
        el = OxmlElement(tag)
        st.append(el)


def code_doc(html_path: Path, out_path: Path):
    text = html_path.read_text(encoding="utf-8")
    lines = text.split("\n")
    doc = Document()
    quiet_settings(doc)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.5)
    sec.top_margin = sec.bottom_margin = Inches(0.5)
    style = doc.styles["Normal"]
    style.font.name = "Consolas"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    style.font.size = Pt(7)
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for p in list(doc.paragraphs):          # drop any template paragraph so the text round-trips exactly
        p._element.getparent().remove(p._element)
    for line in lines:
        p = doc.add_paragraph()
        if line:
            run = p.add_run(line)
            no_proof(run)
    doc.save(out_path)
    # round trip
    back = Document(out_path)
    got = "\n".join(p.text for p in back.paragraphs)
    if got != text:
        # find first difference for the report
        n = next((i for i in range(min(len(got), len(text))) if got[i] != text[i]), min(len(got), len(text)))
        raise SystemExit(f"ROUND-TRIP MISMATCH for {html_path.name} at char {n}: got {got[n:n+40]!r} expected {text[n:n+40]!r}")
    return len(lines), len(text), hashlib.sha256(text.encode("utf-8")).hexdigest()[:12], len(text.encode("utf-8"))


def readme_doc(out_path: Path, stats):
    doc = Document()
    quiet_settings(doc)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(1)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Board Ready — how to install from these Word documents", level=1)
    doc.add_paragraph("Army email blocks .html attachments, so the app travels inside a Word document as plain text. "
                      "Word turns it back into a working file in about a minute. You need Microsoft Word and the Edge or Chrome browser. "
                      "No administrator rights and no internet connection are needed. The Arcade games are inside the same file.")
    doc.add_paragraph("There are three versions. BoardReady-Candidate-Offline.html is for candidates and has no AI connector; "
                      "it is the one to share widely. BoardReady-Candidate.html is the same study app with the "
                      "optional AI connector. BoardReady.html is the full version, which adds the board room tools for board members and "
                      "facilitators. Convert only the one you need.")
    doc.add_heading("Turn a code document back into a file", level=2)
    steps = [
        "Open the Word document for the version you need: “4 - BoardReady-Candidate-Offline.html.docx” to share with candidates, “3 - BoardReady-Candidate.html.docx” for the candidate version with the AI connector, or “2 - BoardReady.html.docx” for the full version. It looks like hundreds of pages of code; that is expected. If a yellow Protected View bar appears, click Enable Editing. Do not change any text.",
        "Choose File, then Save As, then Browse. If you only see Save a Copy, use that.",
        "Pick a folder you can find again, for example a new folder named Board Ready on your Desktop.",
        "Set “Save as type” to “Plain Text (*.txt)”.",
        "In “File name”, type the file name from the table below, including the quotation marks, for example \"BoardReady-Candidate-Offline.html\". Then click Save.",
        "A File Conversion window opens. Choose “Other encoding”, then “Unicode (UTF-8)” in the list. Leave “Insert line breaks” and “Allow character substitution” unchecked. Click OK, then close Word.",
        "Open the folder and double-click the file. It opens in the browser and works with no network. The first time, it offers a short guided tour; the ? button at the top of the page brings the tour back any time.",
    ]
    for s in steps:
        doc.add_paragraph(s, style="List Number")
    doc.add_heading("Check the result", level=2)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Word document", "Save it as", "Size after saving"
    for html_name, docx_name, _who in FILES:
        lines, chars, digest, utf8 = stats[html_name]
        word_bytes = utf8 + (lines - 1) + 3 + 2   # CRLF line ends + UTF-8 byte-order mark + final line end, as measured from Word
        row = t.add_row().cells
        row[0].text, row[1].text, row[2].text = docx_name, html_name, f"{word_bytes:,} bytes (about {word_bytes / 1024 / 1024:.1f} MB)"
    doc.add_paragraph("")
    doc.add_paragraph("To see the exact size, right-click the saved file and choose Properties. A difference of a few bytes is fine; a much smaller file means the save was cut short.")
    doc.add_heading("If something goes wrong", level=2)
    for s in [
        "The file name ends in .html.txt: in File Explorer choose View, Show, File name extensions, then rename the file so it ends in .html.",
        "The browser shows a page of code instead of the app: the file still ends in .txt, or the encoding was not UTF-8. Repeat steps 4 to 6.",
        "Copying all the text into Notepad is very slow or fails for a document this large. Use Save As.",
    ]:
        doc.add_paragraph(s, style="List Bullet")
    doc.add_heading("Keep your notes", level=2)
    doc.add_paragraph("Notes you write in the app are kept in the browser, and some computers clear that at restart. Click Save notes "
                      "(at the top of the page, or in Settings, Your data) and choose this same file in the folder from step 3; replace it "
                      "when asked. Your notes are now inside the file and come back every time you open it. If the browser downloads a copy "
                      "instead, move that copy into the folder and replace the old file.")
    doc.add_heading("Sharing", level=2)
    doc.add_paragraph("Once saved, the file can go in a Teams or SharePoint folder or be emailed inside the Army network, and people just open it. "
                      "Progress stays on each person’s own device. A file with notes saved inside carries those notes, so share a fresh copy.")
    doc.add_paragraph("Training aid, not authoritative guidance. Verify citations, thresholds and local procedures with the Contracting & Policy office before relying on them.")
    doc.save(out_path)


if __name__ == "__main__":
    stats = {}
    for html_name, docx_name, _who in FILES:
        src = DIST / html_name
        if not src.exists():
            raise SystemExit(f"missing {src}; run tools/build.py first")
        stats[html_name] = code_doc(src, OUT / docx_name)
        print(f"{docx_name}: lines={stats[html_name][0]:,} chars={stats[html_name][1]:,} sha={stats[html_name][2]} size={(OUT / docx_name).stat().st_size/1024/1024:.2f} MB  round-trip OK")
    readme_doc(OUT / "1 - READ ME FIRST - How to install Board Ready.docx", stats)
    print("readme written")
